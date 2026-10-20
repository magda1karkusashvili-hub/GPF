import os, threading, hmac
from datetime import date, datetime
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-me")
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bookings.xlsx")
LOCK = threading.Lock()
HEAD = ["ID", "Location", "Date", "Name", "Email", "Phone", "Status", "Created"]
LOCATIONS = ["Gudauri", "Ananuri", "Tbilisi Sea", "Mestia"]
STATUSES = ["NEW", "CONFIRMED", "CANCELLED"]
ADMIN_USER = os.environ.get("ADMIN_USER", "admin")
ADMIN_PASS = os.environ.get("ADMIN_PASS", "")

def safe(v):
    v = str(v).strip()
    return "'" + v if v[:1] in "=+-@" else v

def rows():
    if not os.path.exists(FILE):
        return []
    ws = load_workbook(FILE).active
    return [dict(zip(HEAD, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

def counts(data):
    return {s: sum(r["Status"] == s for r in data) for s in STATUSES}

def add(vals):
    with LOCK:
        if os.path.exists(FILE):
            wb = load_workbook(FILE); ws = wb.active
        else:
            wb = Workbook(); ws = wb.active; ws.append(HEAD)
        ws.append([ws.max_row, *[safe(v) for v in vals], "NEW", datetime.now().strftime("%Y-%m-%d %H:%M")])
        wb.save(FILE)

def set_status(i, s):
    with LOCK:
        wb = load_workbook(FILE); ws = wb.active
        for r in ws.iter_rows(min_row=2):
            if r[0].value == i:
                r[6].value = s
        wb.save(FILE)

def login_required(f):
    @wraps(f)
    def w(*a, **k):
        return f(*a, **k) if session.get("admin") else redirect(url_for("login"))
    return w

@app.route("/")
def index(): return render_template("index.html")
@app.route("/about")
def about(): return render_template("about.html")
@app.route("/gallery")
def gallery(): return render_template("gallery.html")
@app.route("/news")
def news(): return render_template("news.html")
@app.route("/news/fai")
def news_detail(): return render_template("news_detail.html")
@app.route("/document")
def documents(): return render_template("document.html")
@app.route("/contact")
def contact(): return render_template("contact.html")

@app.route("/dashboard")
def dashboard():
    data = rows()
    return render_template("dashboard.html", total=len(data), c=counts(data))

@app.route("/BookYourFlight", methods=["GET", "POST"])
def book():
    if request.method == "POST":
        f = request.form
        loc, d = f.get("location", ""), f.get("date", "")
        name, email, phone = (f.get(k, "").strip() for k in ("name", "email", "phone"))
        try: ok = date.fromisoformat(d) >= date.today()
        except ValueError: ok = False
        if loc in LOCATIONS and ok and name and "@" in email and phone:
            add([loc, d, name[:80], email[:120], phone[:30]])
            flash("Request sent. We will contact you to confirm your flight.")
            return redirect(url_for("book"))
        flash("Please check the form: choose a location and a date from today onwards.")
    return render_template("BookYourFlight.html", locations=LOCATIONS, today=date.today().isoformat())

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u, p = request.form.get("username", ""), request.form.get("password", "")
        if ADMIN_PASS and hmac.compare_digest(u, ADMIN_USER) and hmac.compare_digest(p, ADMIN_PASS):
            session["admin"] = True
            return redirect(url_for("admin"))
        flash("Wrong username or password.")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("index"))

@app.route("/admin")
@login_required
def admin():
    data = rows()
    loc = request.args.get("loc", "").lower(); d = request.args.get("date", ""); st = request.args.get("status", "")
    shown = [r for r in data if loc in r["Location"].lower() and (not d or r["Date"] == d) and (not st or r["Status"] == st)]
    return render_template("admin.html", rows=shown, total=len(data), c=counts(data), statuses=STATUSES)

@app.route("/admin/status/<int:i>/<s>", methods=["POST"])
@login_required
def admin_status(i, s):
    if s in STATUSES: set_status(i, s)
    return redirect(request.referrer or url_for("admin"))

if __name__ == "__main__":
    app.run(debug=True)
